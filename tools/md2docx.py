#!/usr/bin/env python3
"""Convierte los documentos Markdown de docs/ a .docx (Word) para entrega.

Uso:
    python tools/md2docx.py                 # convierte todos los docs/*.md
    python tools/md2docx.py docs/03-....md   # convierte uno

El .md sigue siendo la fuente de trabajo; el .docx se regenera desde él.
Requiere: python-docx  (pip install python-docx)

Soporta un subconjunto de Markdown suficiente para estos documentos:
- Encabezados #..######
- Párrafos, **negrita**, *cursiva*, `código`
- Listas con - / * y listas numeradas
- Citas >
- Bloques de código ``` ```
- Tablas GitHub (| a | b |)
- Reglas horizontales ---
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_runs(paragraph, text: str) -> None:
    """Agrega texto a un párrafo interpretando negrita/cursiva/código inline."""
    text = text.replace("⟨", "\u27e8").replace("⟩", "\u27e9")
    text = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", text)
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator_row(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[\s:\-|]+\|?", line.strip())) and "-" in line


def convert(md_path: Path) -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # Bloque de código
        if line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # cierre ```
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(12)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            continue

        # Tabla
        if line.strip().startswith("|") and i + 1 < n and is_separator_row(lines[i + 1]):
            header = split_table_row(line)
            i += 2  # saltar header y separador
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, cell_text in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = ""
                add_runs(cell.paragraphs[0], cell_text)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(header)):
                    txt = row[j] if j < len(row) else ""
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], txt)
            doc.add_paragraph()
            continue

        stripped = line.strip()

        # Regla horizontal
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # Encabezados
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).replace("⟨", "\u27e8").replace("⟩", "\u27e9"),
                            level=min(level, 4))
            i += 1
            continue

        # Cita
        if stripped.startswith(">"):
            quote_text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, quote_text)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # Lista numerada
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(1))
            i += 1
            continue

        # Lista con viñeta
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue

        # Línea vacía
        if not stripped:
            i += 1
            continue

        # Párrafo normal
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


def main(argv: list[str]) -> int:
    targets = [Path(a) for a in argv[1:]] if len(argv) > 1 else sorted(DOCS.glob("*.md"))
    for md in targets:
        if not md.exists():
            print(f"  no existe: {md}")
            continue
        out = convert(md)
        print(f"  {md.name}  ->  {out.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))

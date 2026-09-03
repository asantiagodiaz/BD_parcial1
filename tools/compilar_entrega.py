#!/usr/bin/env python3
"""Compila el DOCUMENTO FINAL de entrega del parcial en un unico .docx.

Uso:
    python tools/compilar_entrega.py

Sigue el punto 6 del enunciado (Final submission): documento completo con
contexto y supuestos, FR, NFR, modelo E-R, arquitectura backend y frontend,
primera implementacion del backend (5.6), implementacion AWS (5.7) y el
apendice de prompt log.

Genera docs/PARCIAL1-ENTREGA.docx con:
- Portada (con hueco para los integrantes del equipo)
- Tabla de contenido (campo TOC de Word; F9 para actualizar)
- Los documentos 00..07 + 99 concatenados, cada uno en pagina nueva
- Recuadros "[ INSERTAR AQUI: ... ]" para lo que el equipo debe pegar a mano:
  los 4 diagramas, el enlace al repositorio de GitHub, y las capturas de la
  consola de AWS (Glue Data Catalog, RDS cifradas, EC2).
- Las evidencias de la prueba de concurrencia y del ETL/vistas se insertan
  como texto al final de 5.6 y 5.7 respectivamente.

NO incluye las guias internas (08 defensa, 09 despliegue, 10 interaccion).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import md2docx  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
OUT = DOCS / "PARCIAL1-ENTREGA.docx"

# (archivo, nombre del diagrama a insertar tras el H1 | None)
ORDEN = [
    ("00-contexto-supuestos.md", None),
    ("01-requisitos-funcionales.md", None),
    ("02-requisitos-no-funcionales.md", None),
    ("03-modelo-entidad-relacion.md",
     "Diagrama entidad-relacion (notacion Crow's Foot, 5 bandas verticales)"),
    ("04-arquitectura-backend.md",
     "Diagrama de arquitectura del backend (componentes + flujo de reserva con bloqueo)"),
    ("05-arquitectura-frontend.md",
     "Diagrama de arquitectura del frontend (componentes SPA + mapeo a endpoints)"),
    ("06-implementacion-backend.md", None),
    ("07-aws-etl-glue-costos.md",
     "Diagrama de arquitectura AWS (2 BD + ETL + relacion con el Glue Data Catalog)"),
    ("99-prompt-log.md", None),
]

# Evidencias de ejecucion que se anexan tras una seccion concreta.
EVIDENCIAS_TRAS = {
    "06-implementacion-backend.md": [
        ("Evidencia — prueba de concurrencia (capa de servicio, 20 hilos simultaneos)",
         "evidencia/prueba-concurrencia.txt"),
        ("Evidencia — prueba de concurrencia (API HTTP en vivo, 20 peticiones)",
         "evidencia/prueba-concurrencia-http.txt"),
        ("Evidencia — prueba de concurrencia contra la API desplegada en AWS",
         "evidencia/prueba-concurrencia-aws.txt"),
    ],
    "07-aws-etl-glue-costos.md": [
        ("Evidencia — ETL + 3 vistas de negocio + idempotencia (local)",
         "evidencia/etl-vistas-negocio.txt"),
        ("Evidencia — Glue Data Catalog: crawlers y tablas catalogadas (AWS)",
         "evidencia/glue-data-catalog-aws.txt"),
        ("Evidencia — 3 vistas de negocio ejecutadas contra la RDS OLAP (AWS)",
         "evidencia/etl-vistas-negocio-aws.txt"),
    ],
}

# Recuadros de captura extra tras una seccion.
CAPTURAS_TRAS = {
    "06-implementacion-backend.md": [
        "Enlace al repositorio de GitHub con el codigo del backend y el esquema de BD",
        "Captura: Swagger UI de la API (http://<IP-EC2>:8000/docs) con los 3 endpoints",
    ],
    "07-aws-etl-glue-costos.md": [
        "Captura: consola de AWS Glue -> Data Catalog -> Databases -> aerolinea_catalog -> Tables (tablas oltp_* y olap_*)",
        "Captura: consola de RDS mostrando las 2 instancias con Encryption = Enabled (NFR-07)",
        "Captura: consola de EC2 con la instancia aerolinea-api en estado Running",
        "Enlace + captura del estimado en la AWS Pricing Calculator (calculator.aws) con los recursos del stack",
    ],
}

_DRAWIO_RE = re.compile(r"\.drawio")


def _page_break(doc: Document) -> None:
    run = doc.add_paragraph().add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def _add_toc(doc: Document) -> None:
    run = doc.add_paragraph().add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "Actualice la tabla de contenido: clic derecho > Actualizar campos (F9)."
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)


def _recuadro(doc: Document, titulo: str, detalle: str, alto: int = 4) -> None:
    """Recuadro visible con borde para pegar contenido a mano."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = tbl.rows[0].cells[0]
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(f"[ INSERTAR AQUI: {titulo} ]")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
    if detalle:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(detalle)
        r2.italic = True
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    for _ in range(alto):
        cell.add_paragraph()
    doc.add_paragraph()


def _portada(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Sistema de reservas aereas"); r.bold = True; r.font.size = Pt(26)
    s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = s.add_run("Parcial 1 - Big Data e Ingenieria de Datos")
    rs.font.size = Pt(14); rs.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s2.add_run("Documento final (entrega). Despliegue: AWS Academy Learner Lab.").font.size = Pt(11)
    for _ in range(2):
        doc.add_paragraph()
    _recuadro(doc, "Integrantes del equipo",
              "Nombres completos y roles. (Si aplica: la restriccion de la seccion 3 se aclaro que no existe; el profesor hara una pregunta en la sustentacion.)",
              alto=3)
    _page_break(doc)
    doc.add_heading("Tabla de contenido", level=1)
    _add_toc(doc)
    _page_break(doc)


def _append_md(doc: Document, md_path: Path, diagrama: str | None) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    # quitar las citas "> Diagrama editable: ...drawio..."
    limpio = [
        ln for ln in lines
        if not (_DRAWIO_RE.search(ln) and ln.lstrip().startswith(">"))
    ]

    insert_at = None
    for i, ln in enumerate(limpio):
        if ln.startswith("# "):
            insert_at = i + 1
            break

    if diagrama and insert_at is not None:
        _parse_into(doc, limpio[:insert_at])
        _recuadro(doc, "captura del diagrama", diagrama, alto=6)
        _parse_into(doc, limpio[insert_at:])
    else:
        _parse_into(doc, limpio)


def _parse_into(doc: Document, lines: list[str]) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        if line.strip().startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"; run.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(12)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            continue
        if line.strip().startswith("|") and i + 1 < n and md2docx.is_separator_row(lines[i + 1]):
            header = md2docx.split_table_row(line)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(md2docx.split_table_row(lines[i])); i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, cell_text in enumerate(header):
                cell = table.rows[0].cells[j]
                cell.text = ""
                md2docx.add_runs(cell.paragraphs[0], cell_text)
                for rr in cell.paragraphs[0].runs:
                    rr.bold = True
            for row in rows:
                cells = table.add_row().cells
                for j in range(len(header)):
                    cells[j].text = ""
                    md2docx.add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else "")
            doc.add_paragraph()
            continue
        stripped = line.strip()
        if stripped in ("---", "***", "___"):
            doc.add_paragraph(); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
            i += 1; continue
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            md2docx.add_runs(p, stripped.lstrip(">").strip())
            for rr in p.runs:
                rr.italic = True
                rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1; continue
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            md2docx.add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1; continue
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            md2docx.add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1; continue
        if not stripped:
            i += 1; continue
        md2docx.add_runs(doc.add_paragraph(), stripped)
        i += 1


def _anexar_evidencia(doc: Document, titulo: str, rel: str) -> None:
    path = DOCS / rel
    _page_break(doc)
    doc.add_heading(titulo, level=3)
    if not path.exists():
        _recuadro(doc, "pegar la salida de esta evidencia",
                  f"archivo esperado: docs/{rel}", alto=6)
        return
    p = doc.add_paragraph()
    run = p.add_run(path.read_text(encoding="utf-8"))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)


def main() -> int:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _portada(doc)

    for idx, (name, diagrama) in enumerate(ORDEN):
        md = DOCS / name
        if not md.exists():
            print(f"  FALTA: {name}"); continue
        if idx > 0:
            _page_break(doc)
        _append_md(doc, md, diagrama)
        nota = "  (+recuadro diagrama)" if diagrama else ""
        print(f"  + {name}{nota}")

        for titulo in CAPTURAS_TRAS.get(name, []):
            _recuadro(doc, "captura / enlace", titulo, alto=5)
            print(f"      recuadro: {titulo[:60]}...")

        for titulo, rel in EVIDENCIAS_TRAS.get(name, []):
            _anexar_evidencia(doc, titulo, rel)
            print(f"      evidencia: {rel}")

    doc.save(OUT)
    print(f"\n  -> {OUT}")
    print("  Abre en Word, F9 en la TOC, y rellena los recuadros rojos:")
    print("   - integrantes del equipo (portada)")
    print("   - 4 diagramas (5.3, 5.4, 5.5, 5.7)")
    print("   - enlace a GitHub + captura de Swagger (tras 5.6)")
    print("   - 3 capturas de la consola AWS (tras 5.7)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
